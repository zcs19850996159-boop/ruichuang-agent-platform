from __future__ import annotations

import ast
import csv
import json
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "work" / "data_unzip" / "data"
END_DIR = ROOT / "work" / "end_unzip" / "end"
OUTPUT_DIR = ROOT / "outputs" / "rag_assets"

MANUAL_SPECS = {
    "空调手册": ("Manual01", "空调"),
    "人体工学椅手册": ("Manual02", "人体工学椅"),
    "空气净化器手册": ("Manual03", "空气净化器"),
    "吹风机手册": ("Manual04", "吹风机"),
    "蒸汽清洁机手册": ("Manual05", "蒸汽清洁机"),
    "洗碗机手册": ("Manual06", "洗碗机"),
    "电钻手册": ("Manual11", "电钻"),
    "健身单车手册": ("Manual14", "健身单车"),
    "健身追踪器手册": ("Manual16", "健身追踪器"),
    "冰箱手册": ("Manual17", "冰箱"),
    "发电机手册": ("Manual18", "发电机"),
    "功能键盘手册": ("Manual21", "功能键盘"),
    "儿童电动摩托车手册": ("Manual26", "儿童电动摩托车"),
    "蓝牙激光鼠标手册": ("Manual27", "蓝牙激光鼠标"),
    "烤箱手册": ("Manual28", "烤箱"),
    "相机手册": ("Manual29", "相机"),
    "水泵手册": ("Manual31", "水泵"),
    "可编程温控器手册": ("Manual36", "可编程温控器"),
    "VR头显手册": ("Manual38", "VR头显"),
    "摩托艇手册": ("Manual40", "摩托艇"),
    "汇总英文手册": ("EN_SUMMARY", "英文汇总手册"),
}

PREFERRED_MAIN_SHEETS = (
    "可导入映射",
    "PIC校正映射",
    "完整校正表",
    "完整核对表",
    "PIC核对表",
    "PIC映射",
    "PIC映射核对",
    "Full_PIC_Mapping",
)

SUPPLEMENT_SHEETS = (
    "caption约束",
    "额外图片",
    "补充和剔除图片",
    "异常与补图",
    "Risk_Details",
    "Summary",
    "结论",
    "核对结论",
    "总结",
    "校正说明",
    "核对说明",
)


@dataclass
class Manual:
    manual_id: str
    product: str
    source_file: str
    text: str
    image_ids: list[str]
    parse_method: str

    @property
    def pic_count(self) -> int:
        return self.text.count("<PIC>")


def as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def as_int(value: Any) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    text = str(value)
    match = re.search(r"\d+", text)
    return int(match.group()) if match else None


def compact_text(text: str, limit: int = 1200) -> str:
    text = re.sub(r"\s+", " ", as_text(text)).strip()
    if len(text) <= limit:
        return text
    return text[: limit - 1] + "…"


def loose_unescape(value: str) -> str:
    value = value.replace(r"\"", '"')
    value = value.replace(r"\/", "/")
    value = value.replace(r"\n", "\n")
    value = value.replace(r"\r", "\r")
    value = value.replace(r"\t", "\t")

    def decode_unicode(match: re.Match[str]) -> str:
        return chr(int(match.group(1), 16))

    return re.sub(r"\\u([0-9a-fA-F]{4})", decode_unicode, value)


def parse_loose_manual(raw: str) -> tuple[str, list[str]]:
    delimiter = raw.rfind('", [')
    if delimiter == -1:
        delimiter = raw.rfind('",[')
    if delimiter == -1:
        raise ValueError("cannot locate loose manual image list delimiter")
    text_raw = raw[2:delimiter]
    image_raw = raw[delimiter + 2 :].strip()
    if image_raw.endswith("]"):
        image_payload = image_raw[1:] if image_raw.startswith(",") else image_raw
    else:
        raise ValueError("loose manual image list has no closing bracket")
    image_ids = ast.literal_eval(image_payload)
    return loose_unescape(text_raw), [str(item) for item in image_ids]


def parse_manual_file(path: Path) -> Manual:
    name = path.stem
    if name not in MANUAL_SPECS:
        raise ValueError(f"unknown manual file: {path.name}")
    manual_id, product = MANUAL_SPECS[name]
    raw = path.read_text(encoding="utf-8")
    if name == "汇总英文手册":
        texts: list[str] = []
        image_ids: list[str] = []
        loose_count = 0
        for line_no, line in enumerate(raw.splitlines(), start=1):
            line = line.strip()
            if not line:
                continue
            try:
                payload = ast.literal_eval(line)
                text, ids = payload[0], payload[1]
            except Exception:
                loose_count += 1
                try:
                    text, ids = parse_loose_manual(line)
                except Exception as exc:
                    raise ValueError(f"failed to parse English summary line {line_no}: {exc}") from exc
            texts.append(str(text))
            image_ids.extend(str(item) for item in ids)
        return Manual(
            manual_id=manual_id,
            product=product,
            source_file=path.name,
            text="\n\n# --- English manual boundary ---\n\n".join(texts),
            image_ids=image_ids,
            parse_method=f"multi_line:{len(texts)}:loose={loose_count}",
        )
    method = "ast"
    try:
        payload = ast.literal_eval(raw)
        text, image_ids = payload[0], payload[1]
    except Exception:
        method = "loose"
        text, image_ids = parse_loose_manual(raw)
    return Manual(
        manual_id=manual_id,
        product=product,
        source_file=path.name,
        text=str(text),
        image_ids=[str(item) for item in image_ids],
        parse_method=method,
    )


def derive_manual_id_from_file(path: Path) -> str:
    match = re.search(r"Manual\d+", path.name)
    return match.group(0) if match else "EN_SUMMARY"


def product_for_manual_id(manual_id: str) -> str:
    for mid, product in MANUAL_SPECS.values():
        if mid == manual_id:
            return product
    return "英文汇总手册"


def row_values(row) -> list[Any]:
    values = list(row)
    while values and values[-1] is None:
        values.pop()
    return values


def header_score(values: list[Any]) -> int:
    tokens = " ".join(as_text(v) for v in values)
    score = 0
    for token in ("图片编号", "image_id", "Image ID", "校正后图片编号", "PIC序号", "pic_index", "Manual Key"):
        if token in tokens:
            score += 2
    for token in ("caption", "Caption", "人工caption", "图片实际内容", "风险", "status", "状态"):
        if token in tokens:
            score += 1
    if len(values) >= 3:
        score += 1
    return score


def find_header(rows: list[tuple[int, list[Any]]]) -> tuple[int, list[str]] | None:
    candidates = [(idx, vals, header_score(vals)) for idx, vals in rows if header_score(vals) >= 3]
    if not candidates:
        return None
    idx, values, _score = max(candidates, key=lambda item: (item[2], -item[0]))
    return idx, [as_text(v) for v in values]


def rows_from_sheet(ws) -> list[tuple[int, list[Any]]]:
    rows: list[tuple[int, list[Any]]] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = row_values(row)
        if any(v is not None and as_text(v) != "" for v in values):
            rows.append((idx, values))
    return rows


def pick_col(header: list[str], options: tuple[str, ...], contains: bool = False) -> int | None:
    normalized = [h.lower().replace(" ", "") for h in header]
    for option in options:
        opt = option.lower().replace(" ", "")
        for idx, h in enumerate(normalized):
            if (contains and opt in h) or (not contains and h == opt):
                return idx
    return None


def get_value(values: list[Any], idx: int | None) -> Any:
    if idx is None or idx >= len(values):
        return None
    return values[idx]


def classify_sheet(sheet_name: str) -> int:
    if sheet_name in PREFERRED_MAIN_SHEETS:
        return PREFERRED_MAIN_SHEETS.index(sheet_name)
    return 999


def has_image_id_like(value: str) -> bool:
    if not value:
        return False
    if value.lower() in {"无", "none", "null", "nan"}:
        return False
    return bool(re.search(r"(Manual\d+|[A-Za-z]+[_-][A-Za-z0-9]+|\d+)", value))


def parse_main_sheet(path: Path, ws, rows: list[tuple[int, list[Any]]]) -> list[dict[str, Any]]:
    header_info = find_header(rows)
    if not header_info:
        return []
    header_row, header = header_info
    data_rows = [(idx, vals) for idx, vals in rows if idx > header_row]
    default_manual_id = derive_manual_id_from_file(path)

    col_manual = pick_col(header, ("manual_id", "Manual Key", "manual key"), contains=True)
    col_pic = pick_col(
        header,
        ("pic_index", "corrected_pic_index", "校正PIC序号", "PIC序号", "PIC No.", "原始文本PIC序号", "序号"),
        contains=False,
    )
    col_pic_alt = pick_col(header, ("PIC",), contains=False)
    col_image = pick_col(header, ("image_id", "校正后图片编号", "图片编号", "Image ID"), contains=False)
    col_raw = pick_col(header, ("raw_image_id", "原数组图片编号", "原图片编号", "原图片编号"), contains=False)
    col_caption = pick_col(
        header,
        (
            "caption",
            "人工caption",
            "人工 caption",
            "人工caption/实际内容",
            "人工caption/约束",
            "图片实际内容 / 人工caption",
            "图片实际内容 / caption",
            "图片实际内容/Caption",
            "图片实际内容/人工caption",
            "图片实际内容/建议caption",
            "图片实际内容",
            "人工caption / 图像内容",
            "核心caption约束",
            "建议caption/关键词",
            "建议约束",
        ),
        contains=False,
    )
    col_section = pick_col(
        header,
        ("section", "section_title", "对应手册段落", "对应章节", "对应章节/步骤", "原文位置/章节", "章节/位置"),
        contains=False,
    )
    col_status = pick_col(
        header,
        (
            "status",
            "状态",
            "匹配状态",
            "匹配判断",
            "匹配结论",
            "是否匹配",
            "与文本是否匹配",
            "图片状态",
            "是否建议作为<PIC>",
        ),
        contains=False,
    )
    col_action = pick_col(
        header,
        ("action", "建议", "建议调整", "调整建议", "建议操作", "处理建议", "binding_action", "use_policy", "校正动作"),
        contains=False,
    )
    col_risk = pick_col(header, ("risk", "risk_level", "风险", "风险等级"), contains=False)
    col_notes = pick_col(
        header,
        ("notes", "备注", "binding_note", "备注/用于检索的caption建议", "备注/caption约束", "异常说明", "说明"),
        contains=False,
    )
    col_before = pick_col(header, ("前文窗口", "前文片段", "PIC前文", "Before Context"), contains=False)
    col_after = pick_col(header, ("后文窗口", "后文片段", "PIC后文", "After Context"), contains=False)

    records: list[dict[str, Any]] = []
    for row_idx, values in data_rows:
        manual_id = as_text(get_value(values, col_manual)) or default_manual_id
        image_id = as_text(get_value(values, col_image))
        raw_image_id = as_text(get_value(values, col_raw))
        if not image_id and raw_image_id:
            image_id = raw_image_id
        pic_index = as_int(get_value(values, col_pic))
        if pic_index is None:
            pic_index = as_int(get_value(values, col_pic_alt))
        if pic_index is None:
            pic_index = len(records) + 1
        row_text = " | ".join(as_text(v) for v in values if v is not None)
        if not has_image_id_like(image_id):
            if re.search(r"留空|忽略|无图片|未找到|no[_ ]?image|blank|缺位|不强行", row_text, flags=re.I):
                image_id = ""
            else:
                continue
        caption = as_text(get_value(values, col_caption))
        if not caption:
            # If a sheet has no explicit caption but has context columns, use a short context summary.
            caption = compact_text(" ".join(as_text(v) for v in values[2:5] if v is not None), 300)
        before = as_text(get_value(values, col_before))
        after = as_text(get_value(values, col_after))
        records.append(
            {
                "manual_id": manual_id,
                "product": product_for_manual_id(manual_id),
                "pic_index": pic_index,
                "image_id": image_id or None,
                "raw_image_id": raw_image_id or None,
                "caption": caption,
                "section": as_text(get_value(values, col_section)) or None,
                "status": as_text(get_value(values, col_status)) or None,
                "action": as_text(get_value(values, col_action)) or None,
                "risk": as_text(get_value(values, col_risk)) or None,
                "notes": as_text(get_value(values, col_notes)) or None,
                "before_context": compact_text(before, 500) if before else None,
                "after_context": compact_text(after, 500) if after else None,
                "source_workbook": path.name,
                "source_sheet": ws.title,
                "source_row": row_idx,
                "mapping_source": "human_review",
            }
        )
    return records


def parse_supplement_sheet(path: Path, ws, rows: list[tuple[int, list[Any]]]) -> list[dict[str, Any]]:
    header_info = find_header(rows)
    if not header_info:
        # Key-value summary sheets still matter for risk_cases.
        return [
            {
                "case_type": "summary_note",
                "manual_id": derive_manual_id_from_file(path),
                "product": product_for_manual_id(derive_manual_id_from_file(path)),
                "source_workbook": path.name,
                "source_sheet": ws.title,
                "source_row": row_idx,
                "issue": compact_text(" | ".join(as_text(v) for v in values), 500),
            }
            for row_idx, values in rows
            if any(v is not None for v in values)
        ]
    header_row, header = header_info
    data_rows = [(idx, vals) for idx, vals in rows if idx > header_row]
    col_manual = pick_col(header, ("manual_id", "Manual Key", "manual key"), contains=True)
    col_pic = pick_col(header, ("PIC No.", "PIC序号", "图片编号/位置", "原数组位置", "Seq"), contains=False)
    col_image = pick_col(header, ("image_id", "图片编号", "Image ID", "原图片编号"), contains=False)
    col_issue = pick_col(
        header,
        ("Issue", "问题/来源", "现状", "状态", "处理建议", "说明", "Recommended Action", "Recommendation", "Risk Level"),
        contains=False,
    )
    col_caption = pick_col(header, ("建议caption/关键词", "建议约束", "实际内容", "说明", "Recommendation"), contains=False)
    records: list[dict[str, Any]] = []
    for row_idx, values in data_rows:
        text = " | ".join(as_text(v) for v in values if v is not None)
        if not text:
            continue
        manual_id = as_text(get_value(values, col_manual)) or derive_manual_id_from_file(path)
        image_id = as_text(get_value(values, col_image)) or None
        records.append(
            {
                "case_type": "review_note",
                "manual_id": manual_id,
                "product": product_for_manual_id(manual_id),
                "pic_index": as_int(get_value(values, col_pic)),
                "image_id": image_id,
                "issue": compact_text(as_text(get_value(values, col_issue)) or text, 600),
                "caption_hint": compact_text(as_text(get_value(values, col_caption)), 300) or None,
                "source_workbook": path.name,
                "source_sheet": ws.title,
                "source_row": row_idx,
            }
        )
    return records


def parse_review_workbooks() -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    all_mappings: list[dict[str, Any]] = []
    risk_cases: list[dict[str, Any]] = []
    for path in sorted(END_DIR.glob("*.xlsx")):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
        sheet_rows = {ws.title: rows_from_sheet(ws) for ws in wb.worksheets}
        main_candidates: list[tuple[int, str, list[dict[str, Any]]]] = []
        for ws in wb.worksheets:
            rows = sheet_rows[ws.title]
            parsed = parse_main_sheet(path, ws, rows)
            if parsed:
                main_candidates.append((classify_sheet(ws.title), ws.title, parsed))
        if main_candidates:
            main_candidates.sort(key=lambda item: item[0])
            _rank, _sheet_name, chosen = main_candidates[0]
            all_mappings.extend(chosen)
        for ws in wb.worksheets:
            if ws.title in SUPPLEMENT_SHEETS or ws.title not in PREFERRED_MAIN_SHEETS:
                risk_cases.extend(parse_supplement_sheet(path, ws, sheet_rows[ws.title]))
        wb.close()
    return dedupe_mapping_records(all_mappings), risk_cases


def dedupe_mapping_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    deduped: list[dict[str, Any]] = []
    seen: set[tuple[str, int, str, str]] = set()
    for rec in records:
        key = (rec["manual_id"], int(rec["pic_index"]), rec.get("image_id") or "", rec["source_workbook"])
        if key in seen:
            continue
        seen.add(key)
        deduped.append(rec)
    return sorted(deduped, key=lambda r: (r["manual_id"], int(r["pic_index"]), r.get("image_id") or ""))


def build_raw_mapping(manuals: dict[str, Manual]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for manual in manuals.values():
        for idx, image_id in enumerate(manual.image_ids, start=1):
            rows.append(
                {
                    "manual_id": manual.manual_id,
                    "product": manual.product,
                    "pic_index": idx,
                    "image_id": image_id,
                    "raw_image_id": image_id,
                    "caption": "",
                    "section": None,
                    "status": "raw_array",
                    "action": None,
                    "risk": None,
                    "notes": None,
                    "before_context": None,
                    "after_context": None,
                    "source_workbook": None,
                    "source_sheet": None,
                    "source_row": None,
                    "mapping_source": "raw_manual_array",
                }
            )
    return rows


def merge_mappings(manuals: dict[str, Manual], reviewed: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_manual = defaultdict(list)
    for rec in reviewed:
        by_manual[rec["manual_id"]].append(rec)
    raw_by_manual = defaultdict(list)
    for rec in build_raw_mapping(manuals):
        raw_by_manual[rec["manual_id"]].append(rec)

    final: list[dict[str, Any]] = []
    manual_ids = sorted(set(raw_by_manual) | set(by_manual))
    for manual_id in manual_ids:
        chosen = by_manual.get(manual_id) or raw_by_manual.get(manual_id, [])
        for rec in chosen:
            manual = manuals.get(manual_id)
            if manual:
                rec["product"] = manual.product
                rec["source_text_file"] = manual.source_file
            final.append(rec)
    return sorted(final, key=lambda r: (r["manual_id"], int(r["pic_index"]), r.get("image_id") or ""))


def normalize_manual_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(?<!^)(?<!\n)\s+#\s+", "\n# ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def pic_positions(text: str) -> list[tuple[int, int]]:
    return [(m.start(), m.end()) for m in re.finditer(r"<PIC>", text)]


def heading_before(text: str, pos: int) -> str | None:
    prefix = text[:pos]
    matches = list(re.finditer(r"(?m)^#\s*([^\n#]{1,120})", prefix))
    if matches:
        return compact_text(matches[-1].group(1), 120)
    inline_matches = list(re.finditer(r"#\s*([^#\n]{1,120})", prefix[-1200:]))
    return compact_text(inline_matches[-1].group(1), 120) if inline_matches else None


def split_text_chunks(text: str, max_chars: int = 1400, overlap: int = 180) -> list[tuple[int, int, str, str | None]]:
    normalized = normalize_manual_text(text)
    boundaries = [m.start() for m in re.finditer(r"(?m)^#\s+", normalized)]
    if 0 not in boundaries:
        boundaries = [0] + boundaries
    boundaries = sorted(set(boundaries + [len(normalized)]))
    chunks: list[tuple[int, int, str, str | None]] = []
    for start, end in zip(boundaries, boundaries[1:]):
        section = normalized[start:end].strip()
        if not section:
            continue
        title_match = re.match(r"#\s*([^\n]+)", section)
        title = compact_text(title_match.group(1), 120) if title_match else None
        absolute_start = start
        if len(section) <= max_chars:
            chunks.append((absolute_start, absolute_start + len(section), section, title))
            continue
        local_start = 0
        while local_start < len(section):
            local_end = min(len(section), local_start + max_chars)
            chunk_text = section[local_start:local_end].strip()
            if chunk_text:
                chunks.append((absolute_start + local_start, absolute_start + local_end, chunk_text, title))
            if local_end >= len(section):
                break
            local_start = max(0, local_end - overlap)
    return chunks


def mapping_index(records: list[dict[str, Any]]) -> dict[tuple[str, int], list[dict[str, Any]]]:
    index: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for rec in records:
        index[(rec["manual_id"], int(rec["pic_index"]))].append(rec)
    return index


def max_risk(risks: list[str | None]) -> str | None:
    order = {"低": 1, "中": 2, "高": 3}
    chosen = None
    chosen_score = 0
    for risk in risks:
        if not risk:
            continue
        score = order.get(str(risk)[0], 1)
        if score > chosen_score:
            chosen = risk
            chosen_score = score
    return chosen


def build_chunks(manuals: dict[str, Manual], mappings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    idx = mapping_index(mappings)
    chunks: list[dict[str, Any]] = []
    for manual in sorted(manuals.values(), key=lambda m: m.manual_id):
        text = normalize_manual_text(manual.text)
        positions = pic_positions(text)
        for chunk_no, (start, end, chunk_text, title) in enumerate(split_text_chunks(text), start=1):
            pic_indices = [i + 1 for i, (p_start, _p_end) in enumerate(positions) if start <= p_start < end]
            related = [rec for pic_idx in pic_indices for rec in idx.get((manual.manual_id, pic_idx), [])]
            chunks.append(
                {
                    "chunk_id": f"{manual.manual_id}:section:{chunk_no:04d}",
                    "manual_id": manual.manual_id,
                    "product": manual.product,
                    "chunk_type": "section",
                    "section_title": title,
                    "text": chunk_text,
                    "pic_indices": pic_indices,
                    "image_ids": [rec["image_id"] for rec in related if rec.get("image_id")],
                    "captions": [rec["caption"] for rec in related if rec.get("caption")],
                    "risk": max_risk([rec.get("risk") for rec in related]),
                    "source_text_file": manual.source_file,
                }
            )
        for rec in [r for r in mappings if r["manual_id"] == manual.manual_id]:
            pic_index = int(rec["pic_index"])
            image_part = rec.get("image_id") or "NO_IMAGE"
            context = ""
            section = rec.get("section") or None
            if 1 <= pic_index <= len(positions):
                pos = positions[pic_index - 1][0]
                context = text[max(0, pos - 450) : min(len(text), pos + 450)]
                section = section or heading_before(text, pos)
            else:
                context = " ".join(
                    part
                    for part in [
                        as_text(rec.get("before_context")),
                        as_text(rec.get("caption")),
                        as_text(rec.get("after_context")),
                    ]
                    if part
                )
            chunks.append(
                {
                    "chunk_id": f"{manual.manual_id}:pic:{pic_index:04d}:{image_part}",
                    "manual_id": manual.manual_id,
                    "product": manual.product,
                    "chunk_type": "pic_context",
                    "section_title": section,
                    "text": compact_text(context, 1400),
                    "pic_indices": [pic_index],
                    "image_ids": [rec["image_id"]] if rec.get("image_id") else [],
                    "captions": [rec["caption"]] if rec.get("caption") else [],
                    "risk": rec.get("risk"),
                    "source_text_file": manual.source_file,
                }
            )
    return chunks


def automated_risks(manuals: dict[str, Manual], mappings: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_manual = defaultdict(list)
    for rec in mappings:
        by_manual[rec["manual_id"]].append(rec)
    cases: list[dict[str, Any]] = []
    for manual_id, manual in sorted(manuals.items()):
        final_count = len(by_manual.get(manual_id, []))
        raw_duplicates = {k: v for k, v in Counter(manual.image_ids).items() if v > 1}
        final_duplicates = {
            k: v
            for k, v in Counter(r["image_id"] for r in by_manual.get(manual_id, []) if r.get("image_id")).items()
            if v > 1
        }
        if manual.pic_count != len(manual.image_ids):
            cases.append(
                {
                    "case_type": "raw_count_mismatch",
                    "manual_id": manual_id,
                    "product": manual.product,
                    "issue": f"raw text has {manual.pic_count} <PIC> placeholders but image array has {len(manual.image_ids)} ids",
                    "pic_count": manual.pic_count,
                    "image_array_count": len(manual.image_ids),
                    "final_mapping_count": final_count,
                    "source_text_file": manual.source_file,
                }
            )
        if final_count not in {manual.pic_count, len(manual.image_ids)}:
            cases.append(
                {
                    "case_type": "corrected_count_differs_from_raw",
                    "manual_id": manual_id,
                    "product": manual.product,
                    "issue": f"final human-reviewed mapping count is {final_count}; raw placeholders={manual.pic_count}, raw image ids={len(manual.image_ids)}",
                    "pic_count": manual.pic_count,
                    "image_array_count": len(manual.image_ids),
                    "final_mapping_count": final_count,
                    "source_text_file": manual.source_file,
                }
            )
        if raw_duplicates:
            cases.append(
                {
                    "case_type": "raw_duplicate_image_id",
                    "manual_id": manual_id,
                    "product": manual.product,
                    "issue": "raw image array contains duplicate image ids",
                    "duplicates": raw_duplicates,
                    "source_text_file": manual.source_file,
                }
            )
        if final_duplicates:
            cases.append(
                {
                    "case_type": "final_duplicate_image_id",
                    "manual_id": manual_id,
                    "product": manual.product,
                    "issue": "final mapping contains duplicate image ids; this may be intentional reuse",
                    "duplicates": final_duplicates,
                    "source_text_file": manual.source_file,
                }
            )
    return cases


def write_jsonl(path: Path, records: list[dict[str, Any]]) -> None:
    with path.open("w", encoding="utf-8", newline="\n") as f:
        for rec in records:
            f.write(json.dumps(rec, ensure_ascii=False, sort_keys=True) + "\n")


def write_csv(path: Path, records: list[dict[str, Any]], fields: list[str]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        for rec in records:
            writer.writerow(rec)


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    manuals = {manual.manual_id: manual for manual in (parse_manual_file(p) for p in DATA_DIR.glob("*.txt"))}
    reviewed_mappings, review_risks = parse_review_workbooks()
    final_mappings = merge_mappings(manuals, reviewed_mappings)
    chunks = build_chunks(manuals, final_mappings)
    risks = automated_risks(manuals, final_mappings) + review_risks

    manifest = {
        "manual_count": len(manuals),
        "pic_mapping_count": len(final_mappings),
        "manual_chunk_count": len(chunks),
        "risk_case_count": len(risks),
        "manuals": {
            mid: {
                "product": m.product,
                "source_file": m.source_file,
                "parse_method": m.parse_method,
                "text_chars": len(m.text),
                "raw_pic_count": m.pic_count,
                "raw_image_id_count": len(m.image_ids),
                "final_mapping_count": sum(1 for r in final_mappings if r["manual_id"] == mid),
                "section_chunk_count": sum(1 for c in chunks if c["manual_id"] == mid and c["chunk_type"] == "section"),
                "pic_context_chunk_count": sum(1 for c in chunks if c["manual_id"] == mid and c["chunk_type"] == "pic_context"),
            }
            for mid, m in sorted(manuals.items())
        },
    }

    write_jsonl(OUTPUT_DIR / "pic_mapping.jsonl", final_mappings)
    write_jsonl(OUTPUT_DIR / "manual_chunks.jsonl", chunks)
    write_jsonl(OUTPUT_DIR / "risk_cases.jsonl", risks)
    write_csv(
        OUTPUT_DIR / "pic_mapping.csv",
        final_mappings,
        [
            "manual_id",
            "product",
            "pic_index",
            "image_id",
            "raw_image_id",
            "caption",
            "section",
            "status",
            "action",
            "risk",
            "notes",
            "source_workbook",
            "source_sheet",
            "source_row",
            "mapping_source",
        ],
    )
    (OUTPUT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
